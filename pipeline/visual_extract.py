"""
Visual extraction — get pixels out of PDFs, DOCX and raw images.

Pure mechanics, no API calls:
  - per-page PDF text lengths (classification input)
  - PDF page rasterization (pypdfium2 — pip-only, no system deps)
  - embedded-image extraction from PDF pages (pypdf XObjects, size-filtered)
  - embedded-image extraction from DOCX in DOCUMENT ORDER with section anchors
    (a .docx has no pages; the locator is {figure_index, section_title} —
    consistent with how docx text chunks already cite by section)
  - raw image loading incl. HEIC (pillow-heif)
"""
from __future__ import annotations

import io
import logging
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger("pipeline.visual_extract")

# Embedded images smaller than this are logos/bullets, not figures.
MIN_FIGURE_PIXELS = 120 * 120
MIN_FIGURE_BYTES = 6_000


def pdf_page_text_lens(pdf_bytes: bytes) -> List[int]:
    """Characters of extractable text per page (classification input)."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    out = []
    for page in reader.pages:
        try:
            out.append(len((page.extract_text() or "").strip()))
        except Exception:
            out.append(0)
    return out


def pdf_embedded_images(pdf_bytes: bytes, page_index: int) -> List[Tuple[bytes, str]]:
    """
    Embedded raster images on one PDF page, size-filtered.
    Returns [(image_bytes, ext), ...]. Vector line-work is NOT captured here —
    pages that draw their figures as vectors need rasterize_pdf_page().
    """
    from pypdf import PdfReader
    from PIL import Image
    reader = PdfReader(io.BytesIO(pdf_bytes))
    page = reader.pages[page_index]
    out: List[Tuple[bytes, str]] = []
    try:
        images = page.images
    except Exception as e:
        logger.warning("image enumeration failed on page %d: %s", page_index + 1, e)
        return out
    for im in images:
        data = im.data
        if len(data) < MIN_FIGURE_BYTES:
            continue
        try:
            pil = Image.open(io.BytesIO(data))
            if pil.width * pil.height < MIN_FIGURE_PIXELS:
                continue
        except Exception as e:
            logger.debug("unreadable embedded image on page %d skipped: %s",
                         page_index + 1, e)
            continue
        ext = (im.name.rsplit(".", 1)[-1] if "." in im.name else "png").lower()
        out.append((data, ext))
    return out


def rasterize_pdf_page(pdf_bytes: bytes, page_index: int, dpi: int = 150) -> bytes:
    """Render one PDF page to PNG bytes (vector schematics, scanned pages)."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_bytes)
    try:
        page = pdf[page_index]
        bitmap = page.render(scale=dpi / 72.0)
        pil = bitmap.to_pil()
        buf = io.BytesIO()
        pil.save(buf, format="PNG", optimize=True)
        return buf.getvalue()
    finally:
        pdf.close()


def pdf_page_count(pdf_bytes: bytes) -> int:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_bytes)
    try:
        return len(pdf)
    finally:
        pdf.close()


def docx_images(docx_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Embedded images from a .docx in DOCUMENT ORDER, each anchored to the nearest
    preceding heading. Locator = {figure_index, section_title} (docx has no pages).

    Returns [{bytes, ext, figure_index, section_title}, ...].
    """
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(docx_bytes))

    out: List[Dict[str, Any]] = []
    section_title = ""
    fig_index = 0
    body = doc.element.body

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            # Track headings for the section anchor.
            style_el = child.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            if style_el is not None and "Heading" in (style_el.get(qn("w:val")) or ""):
                texts = child.findall(".//" + qn("w:t"))
                title = "".join(t.text or "" for t in texts).strip()
                if title:
                    section_title = title
            # Images anchor via blips inside drawings.
            for blip in child.findall(".//" + qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if not rid:
                    continue
                try:
                    part = doc.part.related_parts[rid]
                except KeyError:
                    continue
                data = part.blob
                if len(data) < MIN_FIGURE_BYTES:
                    continue
                fig_index += 1
                ext = part.partname.ext.lstrip(".").lower() or "png"
                out.append({
                    "bytes": data, "ext": ext,
                    "figure_index": fig_index,
                    "section_title": section_title,
                })
    return out


def load_image_bytes(data: bytes, ext: str) -> Tuple[bytes, str]:
    """
    Normalize any raster input (incl. HEIC) to PNG bytes + media type.
    """
    from PIL import Image
    import pillow_heif
    pillow_heif.register_heif_opener()
    img = Image.open(io.BytesIO(data))
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue(), "image/png"
